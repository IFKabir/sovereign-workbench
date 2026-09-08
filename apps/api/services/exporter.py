"""In-memory Office exports for sovereign workbench responses."""

from __future__ import annotations

import re
import textwrap
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from typing import Literal


ExportFormat = Literal["docx", "pptx", "xlsx"]


@dataclass(slots=True)
class ExportCitation:
    document: str
    section: str = ""
    excerpt: str = ""
    score: float | None = None


@dataclass(slots=True)
class ExportMetric:
    name: str
    value: str


@dataclass(slots=True)
class ExportContent:
    title: str
    content: str
    user_id: str
    role: str
    source: str = "Sovereign AI"
    thread_id: str | None = None
    citations: list[ExportCitation] = field(default_factory=list)
    metrics: list[ExportMetric] = field(default_factory=list)
    calculation_output: str = ""
    calculation_script: str = ""


MIME_TYPES: dict[ExportFormat, str] = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def safe_filename(title: str, export_format: ExportFormat) -> str:
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", title.strip()).strip("-").lower()
    return f"{stem[:72] or 'sovereign-export'}.{export_format}"


def _clean_markdown(text: str) -> str:
    text = re.sub(r"```(?:\w+)?\s*([\s\S]*?)```", r"\1", text)
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"[*_~`]", "", text)
    return text.strip()


def _blocks(text: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw in _clean_markdown(text).splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            blocks.append(("heading3", line[4:]))
        elif line.startswith("## "):
            blocks.append(("heading2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("heading1", line[2:]))
        elif re.match(r"^[-•]\s+", line):
            blocks.append(("bullet", re.sub(r"^[-•]\s+", "", line)))
        elif re.match(r"^\d+[.)]\s+", line):
            blocks.append(("number", re.sub(r"^\d+[.)]\s+", "", line)))
        else:
            blocks.append(("body", line))
    return blocks


def _metadata(export: ExportContent) -> list[tuple[str, str]]:
    return [
        ("Prepared for", export.user_id),
        ("Role", export.role.replace("_", " ").title()),
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Source", export.source),
        ("Thread", export.thread_id or "Not recorded"),
    ]


def build_docx(export: ExportContent) -> bytes:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(export.title)
    intro = document.add_paragraph(
        "This document records the workbench response and its supporting operational context."
    )
    intro.runs[0].italic = True

    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.9)
    for label, value in _metadata(export):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = OxmlElement(f"w:{edge}")
                tag.set(qn("w:val"), "single")
                tag.set(qn("w:sz"), "4")
                tag.set(qn("w:color"), "D9D9D9")
                borders.append(tag)
            tc_pr.append(borders)

    document.add_heading("Response", level=1)
    for kind, value in _blocks(export.content):
        if kind.startswith("heading"):
            document.add_heading(value, level=min(int(kind[-1]), 3))
        elif kind == "bullet":
            document.add_paragraph(value, style="List Bullet")
        elif kind == "number":
            document.add_paragraph(value, style="List Number")
        else:
            document.add_paragraph(value)

    if export.metrics or export.calculation_output:
        document.add_heading("Calculation details", level=1)
        if export.metrics:
            metrics = document.add_table(rows=1, cols=2)
            metrics.style = "Table Grid"
            metrics.rows[0].cells[0].text = "Metric"
            metrics.rows[0].cells[1].text = "Value"
            for cell in metrics.rows[0].cells:
                cell.paragraphs[0].runs[0].bold = True
            for metric in export.metrics:
                cells = metrics.add_row().cells
                cells[0].text, cells[1].text = metric.name, metric.value
        if export.calculation_output:
            document.add_paragraph(_clean_markdown(export.calculation_output))
        if export.calculation_script:
            document.add_heading("Calculation script", level=2)
            paragraph = document.add_paragraph(export.calculation_script)
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)

    if export.citations:
        document.add_heading("References", level=1)
        for citation in export.citations:
            label = citation.document
            if citation.section:
                label += f" — {citation.section}"
            document.add_paragraph(label, style="List Number")
            if citation.excerpt:
                document.add_paragraph(citation.excerpt)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Sovereign AI Workbench | Internal operational use").font.size = Pt(8)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pptx(export: ExportContent) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    accent = RGBColor(87, 105, 44)
    dark = RGBColor(32, 32, 32)

    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = dark
    cover.shapes.title.text = export.title
    cover.placeholders[1].text = f"{export.source}\n{_metadata(export)[2][1]}"
    for shape in cover.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.color.rgb = RGBColor(255, 255, 255)

    blocks = _blocks(export.content)
    sections: list[tuple[str, list[str]]] = []
    current_title = "Response summary"
    current_lines: list[str] = []
    for kind, value in blocks:
        if kind.startswith("heading") and current_lines:
            sections.append((current_title, current_lines))
            current_title, current_lines = value, []
        elif kind.startswith("heading"):
            current_title = value
        else:
            wrapped = textwrap.wrap(value, width=180, break_long_words=False) or [""]
            for line in wrapped:
                current_lines.append(line)
                if len(current_lines) >= 5:
                    sections.append((current_title, current_lines))
                    current_title, current_lines = "Response continued", []
    if current_lines or not sections:
        sections.append((current_title, current_lines or ["No response content was provided."]))

    if export.metrics:
        sections.append(("Calculation results", [f"{m.name}: {m.value}" for m in export.metrics]))
    if export.citations:
        sections.append(("References", [f"{c.document} — {c.section}".rstrip(" —") for c in export.citations]))

    for slide_title, lines in sections:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_title[:100]
        title_frame = slide.shapes.title.text_frame
        title_frame.paragraphs[0].runs[0].font.color.rgb = dark
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, line in enumerate(lines):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(20)
            paragraph.space_after = Pt(10)
        rule = slide.shapes.add_shape(1, Inches(0.65), Inches(1.25), Inches(0.08), Inches(5.4))
        rule.fill.solid()
        rule.fill.fore_color.rgb = accent
        rule.line.fill.background()
        for paragraph in title_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(28)
                run.font.bold = True

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _excel_text(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def build_xlsx(export: ExportContent) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    navy = "233142"
    green = "57692C"
    light = "EAF0E0"
    border = Border(bottom=Side(style="thin", color="D9D9D9"))
    summary.sheet_view.showGridLines = False
    summary["A1"] = _excel_text(export.title)
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color="000000")
    summary.merge_cells("A1:D1")
    summary["A2"] = "Generated workbench response and supporting context"
    summary["A2"].font = Font(name="Arial", size=10, italic=True, color="666666")
    summary.merge_cells("A2:D2")
    row = 4
    for label, value in _metadata(export):
        summary.cell(row, 1, label).font = Font(name="Arial", bold=True)
        summary.cell(row, 2, _excel_text(value))
        row += 1
    row += 1
    summary.cell(row, 1, "Response").font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    summary.cell(row, 1).fill = PatternFill("solid", fgColor=navy)
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    for kind, value in _blocks(export.content):
        chunks = [value[index:index + 30_000] for index in range(0, len(value), 30_000)] or [""]
        for chunk in chunks:
            cell = summary.cell(row, 1, _excel_text(chunk))
            summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Arial", bold=kind.startswith("heading"))
            cell.border = border
            summary.row_dimensions[row].height = min(180, max(18, 15 * (len(chunk) // 110 + 1)))
            row += 1
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 24
    summary.column_dimensions["D"].width = 24
    summary.freeze_panes = "A4"

    if export.metrics or export.calculation_output or export.calculation_script:
        calculations = workbook.create_sheet("Calculations")
        calculations.sheet_view.showGridLines = False
        calculations.append(["Metric", "Value"])
        for metric in export.metrics:
            calculations.append([_excel_text(metric.name), _excel_text(metric.value)])
        if export.calculation_output:
            calculations.append(["Execution output", _excel_text(export.calculation_output)])
        if export.calculation_script:
            calculations.append(["Script", _excel_text(export.calculation_script)])
        calculations.freeze_panes = "A2"
        calculations.column_dimensions["A"].width = 28
        calculations.column_dimensions["B"].width = 90
        for cell in calculations[1]:
            cell.fill = PatternFill("solid", fgColor=green)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        for row_cells in calculations.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    if export.citations:
        refs = workbook.create_sheet("References")
        refs.sheet_view.showGridLines = False
        refs.append(["Document", "Section", "Excerpt", "Relevance score"])
        for citation in export.citations:
            refs.append([
                _excel_text(citation.document),
                _excel_text(citation.section),
                _excel_text(citation.excerpt),
                citation.score,
            ])
        refs.freeze_panes = "A2"
        refs.auto_filter.ref = f"A1:D{refs.max_row}"
        for cell in refs[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        for width, column in zip((28, 36, 80, 18), "ABCD"):
            refs.column_dimensions[column].width = width
        for row_cells in refs.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row_cells in refs.iter_rows(min_row=2, min_col=4, max_col=4):
            row_cells[0].number_format = "0.000"

    for sheet in workbook.worksheets:
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                if cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=10, bold=cell.font.bold)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_export(export: ExportContent, export_format: ExportFormat) -> bytes:
    builders = {"docx": build_docx, "pptx": build_pptx, "xlsx": build_xlsx}
    return builders[export_format](export)
